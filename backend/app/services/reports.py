import io
from calendar import monthrange
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.models.timekeeping import TimesheetRecord, TimesheetStatus
from app.models.user import User

STATUS_LABELS = {
    "present": "Присутствие",
    "absent": "Отсутствие",
    "vacation": "Отпуск",
    "sick": "Больничный",
    "holiday": "Праздник",
    "business_trip": "Командировка",
}

# Standard T-13 codes
T13_CODES = {
    TimesheetStatus.present: "Я",
    TimesheetStatus.absent: "НН",
    TimesheetStatus.vacation: "ОТ",
    TimesheetStatus.sick: "Б",
    TimesheetStatus.holiday: "В",
    TimesheetStatus.business_trip: "К",
}

T13_FILLS = {
    TimesheetStatus.present: "DCFCE7",
    TimesheetStatus.vacation: "DBEAFE",
    TimesheetStatus.sick: "FEF3C7",
    TimesheetStatus.absent: "FEE2E2",
    TimesheetStatus.holiday: "F3F4F6",
    TimesheetStatus.business_trip: "EDE9FE",
}

MONTH_NAMES_RU = [
    "", "январь", "февраль", "март", "апрель", "май", "июнь",
    "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь",
]

_thin = Side(style="thin")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _sc(cell, *, bold=False, size=9, fg="000000", bg=None, align="center", wrap=False):
    cell.font = Font(bold=bold, size=size, color=fg)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    cell.border = _border
    if bg:
        cell.fill = PatternFill("solid", fgColor=bg)


def generate_timesheet_excel(
    db: Session,
    year: int,
    month: int,
    department_id: int | None = None,
    user_id: int | None = None,
) -> bytes:
    num_days = monthrange(year, month)[1]

    query = db.query(TimesheetRecord).join(User, TimesheetRecord.user_id == User.id)
    if department_id:
        query = query.filter(User.department_id == department_id)
    if user_id:
        query = query.filter(TimesheetRecord.user_id == user_id)
    query = query.filter(
        TimesheetRecord.date >= date(year, month, 1),
        TimesheetRecord.date <= date(year, month, num_days),
    )
    records = query.order_by(TimesheetRecord.date).all()

    # Group by employee
    user_day: dict[int, dict[int, TimesheetRecord]] = defaultdict(dict)
    user_obj: dict[int, User] = {}
    for rec in records:
        user_day[rec.user_id][rec.date.day] = rec
        if rec.user:
            user_obj[rec.user_id] = rec.user

    wb = Workbook()
    ws = wb.active
    ws.title = f"Т-13 {month:02d}.{year}"

    # Column layout
    # A(1)=№, B(2)=ФИО, C(3)=Таб.№, D(4)..AH(34)=дни 1-31, AI(35)=дней, AJ(36)=часов
    COL_NO = 1
    COL_NAME = 2
    COL_TAB = 3
    COL_D1 = 4
    COL_TOTAL_D = 35
    COL_TOTAL_H = 36
    LAST_COL = 36

    DARK_BLUE = "1D4ED8"
    MID_BLUE = "3B82F6"

    # ── Title block (rows 1-3) ─────────────────────────────────────────────
    ws.row_dimensions[1].height = 12
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=LAST_COL)
    c = ws.cell(row=1, column=1, value="ООО КЦД")
    c.font = Font(bold=True, size=10)
    c.alignment = Alignment(horizontal="center")

    ws.row_dimensions[2].height = 24
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=LAST_COL)
    c = ws.cell(row=2, column=1, value="ТАБЕЛЬ УЧЁТА РАБОЧЕГО ВРЕМЕНИ")
    c.font = Font(bold=True, size=14)
    c.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[3].height = 13
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=LAST_COL)
    c = ws.cell(
        row=3, column=1,
        value=(
            f"за {MONTH_NAMES_RU[month]} {year} г.   "
            f"Период: 01.{month:02d}.{year} – {num_days:02d}.{month:02d}.{year}"
        ),
    )
    c.font = Font(size=10)
    c.alignment = Alignment(horizontal="center")

    # ── Column headers (rows 4-5) ──────────────────────────────────────────
    ws.row_dimensions[4].height = 36
    ws.row_dimensions[5].height = 20

    # №
    ws.merge_cells(start_row=4, start_column=COL_NO, end_row=5, end_column=COL_NO)
    _sc(ws.cell(row=4, column=COL_NO, value="№"), bold=True, fg="FFFFFF", bg=DARK_BLUE, wrap=True)

    # ФИО
    ws.merge_cells(start_row=4, start_column=COL_NAME, end_row=5, end_column=COL_NAME)
    _sc(
        ws.cell(row=4, column=COL_NAME, value="Фамилия, инициалы\n(должность)"),
        bold=True, fg="FFFFFF", bg=DARK_BLUE, wrap=True,
    )

    # Таб.№
    ws.merge_cells(start_row=4, start_column=COL_TAB, end_row=5, end_column=COL_TAB)
    _sc(ws.cell(row=4, column=COL_TAB, value="Таб.\n№"), bold=True, fg="FFFFFF", bg=DARK_BLUE, wrap=True)

    # Days group header (row 4, spans all 31 day columns)
    ws.merge_cells(start_row=4, start_column=COL_D1, end_row=4, end_column=COL_D1 + 30)
    _sc(
        ws.cell(row=4, column=COL_D1, value="Отметки о явках и неявках по числам месяца"),
        bold=True, fg="FFFFFF", bg=DARK_BLUE,
    )

    # Day numbers (row 5)
    for d in range(1, 32):
        col = COL_D1 + d - 1
        val = d if d <= num_days else ""
        _sc(ws.cell(row=5, column=col, value=val), bold=True, size=8, fg="FFFFFF", bg=MID_BLUE)
        ws.column_dimensions[get_column_letter(col)].width = 3.2

    # Итого group header (row 4)
    ws.merge_cells(start_row=4, start_column=COL_TOTAL_D, end_row=4, end_column=COL_TOTAL_H)
    _sc(ws.cell(row=4, column=COL_TOTAL_D, value="Итого"), bold=True, fg="FFFFFF", bg=DARK_BLUE)

    _sc(ws.cell(row=5, column=COL_TOTAL_D, value="Дней\nявок"), bold=True, size=8, fg="FFFFFF", bg=MID_BLUE, wrap=True)
    _sc(ws.cell(row=5, column=COL_TOTAL_H, value="Часов"), bold=True, size=8, fg="FFFFFF", bg=MID_BLUE)

    # Column widths
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 7
    ws.column_dimensions[get_column_letter(COL_TOTAL_D)].width = 8
    ws.column_dimensions[get_column_letter(COL_TOTAL_H)].width = 8

    # ── Data rows: 2 sub-rows per employee ────────────────────────────────
    row = 6
    for idx, uid in enumerate(sorted(user_day.keys()), 1):
        day_map = user_day[uid]
        uobj = user_obj.get(uid)
        name = uobj.full_name if uobj else f"ID {uid}"
        pos = (uobj.position or "") if uobj else ""
        label = f"{name}\n({pos})" if pos else name

        code_row, hrs_row = row, row + 1
        ws.row_dimensions[code_row].height = 16
        ws.row_dimensions[hrs_row].height = 13

        ws.merge_cells(start_row=code_row, start_column=COL_NO, end_row=hrs_row, end_column=COL_NO)
        _sc(ws.cell(row=code_row, column=COL_NO, value=idx))

        ws.merge_cells(start_row=code_row, start_column=COL_NAME, end_row=hrs_row, end_column=COL_NAME)
        _sc(ws.cell(row=code_row, column=COL_NAME, value=label), align="left", wrap=True)

        ws.merge_cells(start_row=code_row, start_column=COL_TAB, end_row=hrs_row, end_column=COL_TAB)
        _sc(ws.cell(row=code_row, column=COL_TAB, value=str(uid)))

        present_days = 0
        total_hours = 0.0

        for d in range(1, 32):
            col = COL_D1 + d - 1
            cc = ws.cell(row=code_row, column=col)
            hc = ws.cell(row=hrs_row, column=col)
            rec = day_map.get(d)
            if rec and d <= num_days:
                st = rec.status
                cc.value = T13_CODES.get(st, "?")
                _sc(cc, size=8, bg=T13_FILLS.get(st))
                if st == TimesheetStatus.present:
                    hrs = rec.work_hours or 0.0
                    hc.value = hrs if hrs else ""
                    present_days += 1
                    total_hours += hrs
                else:
                    _sc(hc, size=7)
            else:
                _sc(cc, size=8)
                _sc(hc, size=7)

        ws.merge_cells(start_row=code_row, start_column=COL_TOTAL_D, end_row=hrs_row, end_column=COL_TOTAL_D)
        _sc(ws.cell(row=code_row, column=COL_TOTAL_D, value=present_days), bold=True)

        ws.merge_cells(start_row=code_row, start_column=COL_TOTAL_H, end_row=hrs_row, end_column=COL_TOTAL_H)
        _sc(ws.cell(row=code_row, column=COL_TOTAL_H, value=round(total_hours, 1)), bold=True)

        row += 2

    # ── Grand totals footer ────────────────────────────────────────────────
    if row > 6:
        ws.row_dimensions[row].height = 18
        ws.merge_cells(start_row=row, start_column=COL_NO, end_row=row, end_column=COL_TAB)
        _sc(ws.cell(row=row, column=COL_NO, value="Итого по подразделению:"), bold=True, align="left")

        grand_days = sum(
            sum(1 for r in dm.values() if r.status == TimesheetStatus.present)
            for dm in user_day.values()
        )
        grand_hours = sum(
            sum(r.work_hours or 0 for r in dm.values() if r.status == TimesheetStatus.present)
            for dm in user_day.values()
        )
        for d in range(1, 32):
            _sc(ws.cell(row=row, column=COL_D1 + d - 1), size=8)

        _sc(ws.cell(row=row, column=COL_TOTAL_D, value=grand_days), bold=True)
        _sc(ws.cell(row=row, column=COL_TOTAL_H, value=round(grand_hours, 1)), bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_timesheet_pdf(
    db: Session,
    year: int,
    month: int,
    department_id: int | None = None,
    user_id: int | None = None,
) -> bytes:
    query = db.query(TimesheetRecord).join(User, TimesheetRecord.user_id == User.id)
    if department_id:
        query = query.filter(User.department_id == department_id)
    if user_id:
        query = query.filter(TimesheetRecord.user_id == user_id)
    num_days = monthrange(year, month)[1]
    query = query.filter(
        TimesheetRecord.date >= date(year, month, 1),
        TimesheetRecord.date <= date(year, month, num_days),
    )
    records = query.order_by(TimesheetRecord.date).all()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1 * cm, rightMargin=1 * cm)
    styles = getSampleStyleSheet()
    elements = []

    title = Paragraph(
        f"<b>Табель учёта рабочего времени — {month:02d}/{year}</b>", styles["Title"]
    )
    elements.append(title)
    elements.append(Spacer(1, 0.5 * cm))

    table_data = [["ФИО", "Дата", "Приход", "Уход", "Статус", "Часов"]]
    for rec in records:
        uobj = rec.user
        table_data.append([
            uobj.full_name if uobj else "—",
            str(rec.date),
            str(rec.check_in) if rec.check_in else "",
            str(rec.check_out) if rec.check_out else "",
            STATUS_LABELS.get(rec.status.value, rec.status.value),
            str(rec.work_hours or ""),
        ])

    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF6FF")]),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("PADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(t)

    doc.build(elements)
    return buf.getvalue()
